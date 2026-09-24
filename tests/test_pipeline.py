"""Tests run without Ollama: fake embeddings and a scripted LLM stand in."""
import docx
import pytest
from langchain_chroma import Chroma
from langchain_core.embeddings import DeterministicFakeEmbedding
from langchain_core.language_models import FakeListChatModel

from docqa.generate import NOT_FOUND, Answerer
from docqa.ingest import ingest
from docqa.retrieve import HybridRetriever, reciprocal_rank_fusion


@pytest.fixture
def store(tmp_path):
    return Chroma(collection_name="test_docs", embedding_function=DeterministicFakeEmbedding(size=64),
                  persist_directory=str(tmp_path / "db"))


@pytest.fixture
def docs_dir(tmp_path):
    d = tmp_path / "docs"
    d.mkdir()
    (d / "policy.md").write_text("Refunds are issued within 14 days of purchase.\n\n"
                                 "Part number ZX-4410 is excluded from coverage.")
    (d / "faq.txt").write_text("Our office is open Monday to Friday, 9am to 5pm.")
    word = docx.Document()
    word.add_paragraph("Warranty claims require the original invoice.")
    word.save(d / "guide.docx")
    return d


def test_ingest_reads_all_formats_and_is_idempotent(store, docs_dir):
    n = ingest([str(docs_dir)], store, chunk_size=80, chunk_overlap=10)
    assert n >= 3
    ingest([str(docs_dir)], store, chunk_size=80, chunk_overlap=10)  # re-run
    assert len(store.get()["ids"]) == n  # same IDs overwrite, no duplicates
    sources = {m["source"] for m in store.get()["metadatas"]}
    assert sources == {"policy.md", "faq.txt", "guide.docx"}


def test_keyword_search_finds_exact_part_number(store, docs_dir):
    ingest([str(docs_dir)], store, chunk_size=80, chunk_overlap=10)
    hits = HybridRetriever(store, candidate_k=5, top_k=2).retrieve("ZX-4410")
    assert "ZX-4410" in hits[0].page_content


def test_rrf_rewards_documents_ranked_well_in_both_lists():
    from langchain_core.documents import Document
    a, b, c = (Document(page_content=x, metadata={"chunk_id": x}) for x in "abc")
    fused = reciprocal_rank_fusion([[a, b, c], [b, c, a]])
    assert fused[0] is b  # 2nd + 1st beats 1st + 3rd


def test_answer_keeps_only_valid_citations(store, docs_dir):
    ingest([str(docs_dir)], store, chunk_size=80, chunk_overlap=10)
    llm = FakeListChatModel(responses=["Refunds take 14 days [1]. Also see [9]."])
    result = Answerer(llm, HybridRetriever(store, 5, 2)).ask("How long do refunds take?")
    assert [s["n"] for s in result["sources"]] == [1]  # [9] was invented, dropped


def test_empty_index_returns_not_found(store):
    llm = FakeListChatModel(responses=["should not be called"])
    result = Answerer(llm, HybridRetriever(store)).ask("anything")
    assert result["answer"] == NOT_FOUND and result["sources"] == []
